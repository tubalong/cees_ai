from __future__ import annotations

import base64
import zipfile
from datetime import date
from io import BytesIO

import pymupdf
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.core.config import ExtractionConfig, ModelCatalog, Settings
from app.core.runtime import AppRuntime
from app.extraction import UnsupportedFileTypeError, extract_text
from app.main import create_app
from tests.helpers import profile

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---- unit: plain text formats ----


def test_extract_plain_text_utf8() -> None:
    result = extract_text(content_type="text/plain", data="你好，世界\nhello".encode())
    assert result.text == "你好，世界\nhello"
    assert result.engine == "utf-8"


def test_extract_plain_text_falls_back_to_gb18030() -> None:
    result = extract_text(content_type="text/plain", data="中文".encode("gb18030"))
    assert result.text == "中文"


def test_extract_csv_json_markdown() -> None:
    for content_type, value in [
        ("text/csv", "a,b\n1,2\n"),
        ("application/json", '{"a": 1}'),
        ("text/markdown", "# Title"),
    ]:
        result = extract_text(content_type=content_type, data=value.encode("utf-8"))
        assert result.text == value


# ---- unit: office formats ----


def _docx_bytes() -> bytes:
    document = WordDocument()
    document.add_paragraph("Hello paragraph")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_extract_docx_text() -> None:
    result = extract_text(content_type=DOCX_CONTENT_TYPE, data=_docx_bytes())
    assert "Hello paragraph" in result.text
    assert "A | B" in result.text
    assert result.engine == "python-docx"


def _pptx_bytes() -> bytes:
    slide_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
       xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld>
    <p:spTree>
      <p:sp>
        <p:txBody>
          <a:p><a:r><a:t>Slide one text</a:t></a:r></a:p>
        </p:txBody>
      </p:sp>
    </p:spTree>
  </p:cSld>
</p:sld>"""
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("ppt/slides/slide1.xml", slide_xml)
        slide2 = slide_xml.replace("Slide one text", "Slide two text")
        archive.writestr("ppt/slides/slide2.xml", slide2)
    return buffer.getvalue()


def test_extract_pptx_text() -> None:
    result = extract_text(content_type=PPTX_CONTENT_TYPE, data=_pptx_bytes())
    assert "[Slide 1]" in result.text
    assert "Slide one text" in result.text
    assert "[Slide 2]" in result.text
    assert "Slide two text" in result.text
    assert result.engine == "stdlib-zip-xml"


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Hello", "World"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_extract_xlsx_text() -> None:
    result = extract_text(content_type=XLSX_CONTENT_TYPE, data=_xlsx_bytes())
    assert "[Sheet: Sheet]" in result.text
    assert "Hello\tWorld" in result.text
    assert result.engine == "openpyxl"


def test_extract_xlsx_preserves_empty_columns_and_formats_dates() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "收支台账"
    sheet.append(["发生日期", "方向", "金额", "币种", "部门ID", "项目ID", "凭证号"])
    sheet.append([date(2026, 1, 5), "支出", 15000, "CNY", "D003", None, "V2026010501"])
    sheet["A2"].number_format = "yyyy-mm-dd"
    sheet["C2"].number_format = "#,##0.00"
    buffer = BytesIO()
    workbook.save(buffer)

    result = extract_text(content_type=XLSX_CONTENT_TYPE, data=buffer.getvalue())

    assert "2026-01-05\t支出\t15000\tCNY\tD003\t\tV2026010501" in result.text


# ---- unit: best-effort PDF ----


def _pdf_bytes() -> bytes:
    document = pymupdf.open()
    try:
        page = document.new_page()
        page.insert_text((72, 72), "Hello PDF")
        return document.tobytes()
    finally:
        document.close()


def test_extract_pdf_text() -> None:
    result = extract_text(content_type="application/pdf", data=_pdf_bytes())
    assert "Hello PDF" in result.text
    assert result.engine == "pymupdf"


# ---- unit: dispatch and errors ----


def test_extract_unsupported_content_type_raises() -> None:
    try:
        extract_text(content_type="application/vnd.ms-excel", data=b"x")
    except UnsupportedFileTypeError as exc:
        assert exc.content_type == "application/vnd.ms-excel"
    else:
        raise AssertionError("expected UnsupportedFileTypeError")


def test_extract_infers_from_filename_for_generic_content_type() -> None:
    result = extract_text(
        content_type="application/octet-stream",
        data=b"hello",
        filename="notes.txt",
    )
    assert result.text == "hello"


# ---- API ----


def _client() -> TestClient:
    runtime = AppRuntime(
        settings=Settings(node_env="test", ai_internal_token="secret"),
        catalog=None,
        router=None,
        readiness_errors=[],
    )
    return TestClient(create_app(runtime=runtime))


def _payload(data: bytes, content_type: str = DOCX_CONTENT_TYPE) -> dict[str, object]:
    return {
        "request_id": "req-files-extract-test-1",
        "tenant_id": "tenant-1",
        "user_id": "user-1",
        "filename": "brief.docx",
        "content_type": content_type,
        "data_base64": base64.b64encode(data).decode("ascii"),
    }


def test_extract_endpoint_returns_text_parts() -> None:
    with _client() as client:
        response = client.post(
            "/internal/v1/files/extract",
            json=_payload(_docx_bytes()),
            headers={"X-AI-Internal-Token": "secret"},
        )
    assert response.status_code == 200
    body = response.json()
    assert body["request_id"] == "req-files-extract-test-1"
    assert body["metadata"]["engine"] == "python-docx"
    assert body["metadata"]["text_length"] > 0
    assert any("Hello paragraph" in part["text"] for part in body["parts"])


def test_extract_endpoint_requires_token() -> None:
    with _client() as client:
        response = client.post(
            "/internal/v1/files/extract",
            json=_payload(_docx_bytes()),
        )
    assert response.status_code == 401


def test_extract_endpoint_rejects_unsupported_type() -> None:
    with _client() as client:
        response = client.post(
            "/internal/v1/files/extract",
            json=_payload(b"x", content_type="application/vnd.ms-excel"),
            headers={"X-AI-Internal-Token": "secret"},
        )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "UNSUPPORTED_FILE_TYPE"


def test_extract_endpoint_rejects_bad_base64() -> None:
    payload = _payload(b"x")
    payload["data_base64"] = "!!!not-base64!!!"
    with _client() as client:
        response = client.post(
            "/internal/v1/files/extract",
            json=payload,
            headers={"X-AI-Internal-Token": "secret"},
        )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "INVALID_FILE_EXTRACTION_REQUEST"

def test_extract_endpoint_respects_configured_max_bytes() -> None:
    model_catalog = ModelCatalog(
        profiles={"mock": profile(provider="mock")},
        roles={},
        extraction=ExtractionConfig(max_bytes=4),
    )
    runtime = AppRuntime(
        settings=Settings(node_env="test", ai_internal_token="secret"),
        catalog=model_catalog,
        router=None,
        readiness_errors=[],
    )
    with TestClient(create_app(runtime=runtime)) as client:
        response = client.post(
            "/internal/v1/files/extract",
            json=_payload(b"hello", content_type="text/plain"),
            headers={"X-AI-Internal-Token": "secret"},
        )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "FILE_TOO_LARGE"

from __future__ import annotations

from io import BytesIO

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor

from app.api.generated.models import DocumentOptions, PptxSpec
from app.core.errors import AIServiceError
from app.documents.pptx_renderer import (
    _BLOCK_GAP,
    PptxRenderer,
    _measure_block,
    _paginate,
)

_BRAND = RGBColor(0x56, 0x5C, 0xF6)
_BODY_WIDTH_IN = 13.333 - 2 * 0.72


def pptx_data() -> dict[str, object]:
    return {
        "schema_version": "1.0",
        "title": "季度汇报",
        "subtitle": "经营概览",
        "theme": "brand",
        "slides": [
            {
                "title": "关键结论",
                "layout": "title_and_content",
                "blocks": [
                    {"type": "paragraph", "text": "本季度交付达成。"},
                    {"type": "bullet_list", "items": ["营收增长", "成本受控"]},
                    {
                        "type": "table",
                        "columns": ["指标", "数值"],
                        "rows": [["营收", "120%"]],
                    },
                    {"type": "quote", "text": "稳中求进。", "attribution": "管理层"},
                ],
            }
        ],
    }


def render_options() -> DocumentOptions:
    return DocumentOptions.model_validate(
        {"locale": "zh-CN", "template_id": "business-standard", "include_toc": False}
    )


def paragraph_blocks(*texts: str) -> list:
    """按渲染器接受的模型构造纯段落块，避免测试手工拼 pydantic 实例。"""
    spec = PptxSpec.model_validate(
        {
            "schema_version": "1.0",
            "title": "层级测试",
            "slides": [
                {
                    "title": "层级",
                    "layout": "title_and_content",
                    "blocks": [{"type": "paragraph", "text": text} for text in texts],
                }
            ],
        }
    )
    return list(spec.slides[0].blocks)


def shape_with_text(presentation, expected: str):
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text == expected:
                return shape
    raise AssertionError(f"未找到文本为 {expected!r} 的形状")


def first_run(presentation, text: str):
    return shape_with_text(presentation, text).text_frame.paragraphs[0].runs[0]


def test_renders_openable_pptx() -> None:
    rendered = PptxRenderer().render(
        PptxSpec.model_validate(pptx_data()),
        render_options(),
        request_id="req-render-pptx-1",
    )

    assert rendered.content.startswith(b"PK")
    assert rendered.filename.endswith(".pptx")

    presentation = Presentation(BytesIO(rendered.content))
    # 封面 + 内容页
    assert len(presentation.slides) == 2
    text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
    assert "季度汇报" in text
    assert "关键结论" in text
    assert "营收增长" in text


def test_paragraph_subheading_is_emphasised_over_body() -> None:
    """短句段落按小标题排版：加粗、品牌色、字号大于正文，且正文不低于可读下限。"""
    body_text = "定义：AI Agent 是以大语言模型为推理内核的软件系统。"
    rendered = PptxRenderer().render(
        PptxSpec.model_validate(
            {
                "schema_version": "1.0",
                "title": "Agent 讲义",
                "slides": [
                    {
                        "title": "核心概念",
                        "layout": "title_and_content",
                        "blocks": [
                            {"type": "paragraph", "text": "核心特征"},
                            {"type": "paragraph", "text": body_text},
                        ],
                    }
                ],
            }
        ),
        render_options(),
        request_id="req-render-pptx-subheading",
    )
    presentation = Presentation(BytesIO(rendered.content))

    subheading_run = first_run(presentation, "核心特征")
    body_run = first_run(presentation, body_text)

    assert subheading_run.font.bold is True
    assert subheading_run.font.color.rgb == _BRAND
    assert subheading_run.font.size.pt > body_run.font.size.pt
    # 字号下限：投影上小于 14pt 的正文等于不可读，宁可多分页也不缩到这里以下。
    assert body_run.font.size.pt >= 14


def test_paginate_moves_trailing_subheading_to_next_page() -> None:
    """页尾若只剩一个小标题，把它挪到下一页，避免孤立标题。"""
    body, subheading = paragraph_blocks("正" * 80, "核心特征")
    font_pt = 18
    # 高度刚好只容得下「正文 + 小标题」，第三个块必须翻页。
    height = (
        _measure_block(body, font_pt, _BODY_WIDTH_IN)
        + _BLOCK_GAP
        + _measure_block(subheading, font_pt, _BODY_WIDTH_IN)
        + 0.01
    )

    pages = _paginate([body, subheading, body], _BODY_WIDTH_IN, height, font_pt)

    assert pages == [[body], [subheading, body]]


def test_paginate_keeps_subheading_when_next_page_cannot_hold_it() -> None:
    """下一页放不下时保留原位置，避免为了消除孤立标题而制造新的溢出。"""
    body, subheading, long_body = paragraph_blocks("正" * 80, "核心特征", "正" * 900)
    font_pt = 18
    height = (
        _measure_block(body, font_pt, _BODY_WIDTH_IN)
        + _BLOCK_GAP
        + _measure_block(subheading, font_pt, _BODY_WIDTH_IN)
        + 0.01
    )

    pages = _paginate([body, subheading, long_body], _BODY_WIDTH_IN, height, font_pt)

    assert pages == [[body, subheading], [long_body]]


def test_rejects_table_rows_with_wrong_column_count() -> None:
    payload = pptx_data()
    payload["slides"][0]["blocks"][2]["rows"] = [["营收"]]

    with pytest.raises(AIServiceError) as raised:
        PptxRenderer().render(
            PptxSpec.model_validate(payload),
            render_options(),
            request_id="req-render-pptx-2",
        )
    assert raised.value.status_code == 422

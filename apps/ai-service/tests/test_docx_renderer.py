from __future__ import annotations

import base64
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document as WordDocument
from PIL import Image

from app.api.generated.models import DocumentOptions, DocumentSpec
from app.core.errors import AIServiceError
from app.documents.docx_renderer import DocxRenderer, content_disposition


def document_data() -> dict[str, object]:
    return {
        "schema_version": "1.0",
        "title": 'Implementation: Plan/2026?',
        "subtitle": "MVP",
        "sections": [
            {
                "heading": "Scope",
                "level": 1,
                "blocks": [
                    {"type": "paragraph", "text": "Ship the first release."},
                    {"type": "bullet_list", "items": ["Compose", "Render"]},
                    {"type": "numbered_list", "items": ["Review", "Publish"]},
                    {
                        "type": "table",
                        "columns": ["Stage", "Owner"],
                        "rows": [["Build", "Team"]],
                    },
                    {"type": "quote", "text": "Keep it controlled.", "attribution": "CEES"},
                    {"type": "page_break"},
                ],
            }
        ],
        "source_refs": ["requirements"],
    }


def render_options(*, include_toc: bool = True) -> DocumentOptions:
    return DocumentOptions.model_validate(
        {
            "locale": "en-US",
            "template_id": "business-standard",
            "include_toc": include_toc,
        }
    )


def test_renders_openable_docx_with_supported_blocks() -> None:
    rendered = DocxRenderer().render(
        DocumentSpec.model_validate(document_data()),
        render_options(),
        request_id="req-render-1",
    )

    assert rendered.content.startswith(b"PK")
    assert rendered.filename == "Implementation_ Plan_2026_.docx"
    document = WordDocument(BytesIO(rendered.content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert "Implementation: Plan/2026?" in paragraphs
    assert "Scope" in paragraphs
    assert "Ship the first release." in paragraphs
    assert "Compose" in paragraphs
    assert "Review" in paragraphs
    assert "Keep it controlled." in paragraphs
    assert document.tables[0].cell(0, 0).text == "Stage"
    assert document.tables[0].cell(1, 1).text == "Team"
    assert "TOC" in document._element.xml
    assert "PAGE" in document.sections[0].footer._element.xml


def test_rejects_table_rows_with_wrong_column_count() -> None:
    payload = document_data()
    payload["sections"][0]["blocks"][3]["rows"] = [["Build"]]

    with pytest.raises(AIServiceError) as raised:
        DocxRenderer().render(
            DocumentSpec.model_validate(payload),
            render_options(include_toc=False),
            request_id="req-render-2",
        )

    assert raised.value.code == "DOCUMENT_SPEC_INVALID"
    assert raised.value.status_code == 422


def test_rejects_duplicate_source_references() -> None:
    payload = document_data()
    payload["source_refs"] = ["requirements", "requirements"]

    with pytest.raises(AIServiceError) as raised:
        DocxRenderer().render(
            DocumentSpec.model_validate(payload),
            render_options(include_toc=False),
            request_id="req-render-3",
        )

    assert raised.value.code == "DOCUMENT_SPEC_INVALID"


def test_content_disposition_has_ascii_fallback_and_utf8_filename() -> None:
    header = content_disposition("项目方案.docx")

    assert 'filename="document.docx"' in header
    assert "filename*=UTF-8''%E9%A1%B9%E7%9B%AE%E6%96%B9%E6%A1%88.docx" in header


def test_embeds_inline_image_bytes_in_docx_media() -> None:
    image_stream = BytesIO()
    Image.new("RGB", (24, 16), "green").save(image_stream, format="PNG")
    image_bytes = image_stream.getvalue()
    payload = document_data()
    payload["sections"][0]["blocks"].append({
        "type": "image",
        "url": "data:image/png;base64," + base64.b64encode(image_bytes).decode("ascii"),
        "alt": "Team collaboration",
        "caption": "Team collaboration",
    })

    rendered = DocxRenderer().render(
        DocumentSpec.model_validate(payload),
        render_options(include_toc=False),
        request_id="req-render-image",
    )

    with ZipFile(BytesIO(rendered.content)) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        assert len(media) == 1
        assert archive.read(media[0]) == image_bytes
    document = WordDocument(BytesIO(rendered.content))
    assert len(document.inline_shapes) == 1
    assert not any("[" in paragraph.text for paragraph in document.paragraphs)
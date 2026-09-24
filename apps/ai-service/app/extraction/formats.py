from __future__ import annotations

import re
import zipfile
from datetime import date, datetime
from io import BytesIO
from xml.etree import ElementTree

from docx import Document
from openpyxl import load_workbook

DRAWINGML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
SPREADSHEETML_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

_SLIDE_RE = re.compile(r"^ppt/slides/slide(\d+)\.xml$")
_SHEET_RE = re.compile(r"^xl/worksheets/sheet(\d+)\.xml$")


def extract_plain_text(data: bytes) -> str:
    """Decode UTF-8/UTF-16 text and fall back to GB18030 for legacy Chinese text."""
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("gb18030", errors="replace")


def extract_docx_text(data: bytes) -> str:
    """Extract paragraph and table text from a DOCX using python-docx."""
    document = Document(BytesIO(data))
    blocks: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text for cell in row.cells]
            blocks.append(" | ".join(cells))
    return "\n".join(block for block in blocks if block.strip())


def extract_pptx_text(data: bytes) -> str:
    """Extract slide text from a PPTX using stdlib ZIP/XML without new dependencies."""
    with zipfile.ZipFile(BytesIO(data)) as archive:
        names = sorted(
            (name for name in archive.namelist() if _SLIDE_RE.match(name)),
            key=lambda name: int(_SLIDE_RE.match(name).group(1)),
        )
        sections: list[str] = []
        for index, name in enumerate(names, start=1):
            root = ElementTree.fromstring(archive.read(name))
            lines = _paragraph_lines(root, DRAWINGML_NS)
            content = "\n".join(lines)
            sections.append(f"[Slide {index}]\n{content}")
        return "\n\n".join(sections)


def extract_xlsx_text(data: bytes) -> str:
    """Extract stable tabular text while preserving empty columns and Excel dates."""
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        sections: list[str] = []
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for values in sheet.iter_rows(values_only=True):
                cells = [_spreadsheet_value(value) for value in values]
                while cells and cells[-1] == "":
                    cells.pop()
                if any(cells):
                    rows.append("\t".join(cells))
            if rows:
                sections.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
        return "\n\n".join(sections)
    finally:
        workbook.close()


def _spreadsheet_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time() == datetime.min.time():
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _paragraph_lines(root: ElementTree.Element, namespace: str) -> list[str]:
    lines: list[str] = []
    for paragraph in root.iter(f"{{{namespace}}}p"):
        text = "".join(paragraph.itertext()).strip()
        if text:
            lines.append(text)
    return lines


def _shared_strings(archive: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in archive.namelist():
        return []
    root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
    return ["".join(item.itertext()) for item in root.iter(f"{{{SPREADSHEETML_NS}}}si")]


def _sheet_rows(root: ElementTree.Element, shared: list[str]) -> list[str]:
    rows: list[str] = []
    for row in root.iter(f"{{{SPREADSHEETML_NS}}}row"):
        cells: list[str] = []
        for cell in row.iter(f"{{{SPREADSHEETML_NS}}}c"):
            value = _cell_value(cell, shared)
            if value != "":
                cells.append(value)
        if cells:
            rows.append("\t".join(cells))
    return rows


def _cell_value(cell: ElementTree.Element, shared: list[str]) -> str:
    kind = cell.get("t")
    if kind == "inlineStr":
        return "".join(cell.itertext()).strip()
    value_element = cell.find(f"{{{SPREADSHEETML_NS}}}v")
    if kind == "s" and value_element is not None and value_element.text:
        try:
            index = int(value_element.text)
            return shared[index] if 0 <= index < len(shared) else ""
        except ValueError:
            return ""
    if value_element is not None and value_element.text:
        return value_element.text
    return ""

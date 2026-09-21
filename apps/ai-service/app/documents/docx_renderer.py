from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from typing import Any
from urllib.parse import quote

from docx import Document as WordDocument
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.api.generated.models import (
    BulletListBlock,
    DocumentOptions,
    DocumentSpec,
    ImageBlock,
    NumberedListBlock,
    PageBreakBlock,
    ParagraphBlock,
    QuoteBlock,
    TableBlock,
    TemplateId,
)
from app.documents.images import load_image
from app.documents.normalize import normalize_document_spec
from app.documents.validation import validate_document_spec

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_INVALID_FILENAME = re.compile(r'[\x00-\x1f<>:"/\\|?*]+')

# 与 `_configure_document` 的页面设置保持一致：A4 宽 21cm，左右页边距各 3.18cm。
_CONTENT_WIDTH_CM = 21.0 - 3.18 * 2
_MAX_IMAGE_HEIGHT_CM = 19.5


@dataclass(frozen=True)
class RenderedDocx:
    content: bytes
    filename: str


class DocxRenderer:
    def render(
        self,
        document: DocumentSpec,
        options: DocumentOptions,
        *,
        request_id: str,
    ) -> RenderedDocx:
        # 渲染前统一清洗：合并「封面」小节、剥离「标题：」等标签前缀、剔除占位行，
        # 保证 DOCX 与 PPTX/PDF 的成稿一致，不出现重复封面与占位符。
        document = normalize_document_spec(document)
        validate_document_spec(document, request_id=request_id, status_code=422)
        # DOCX 暂时共用稳定的 Word 样式；接受全部契约模板 ID，避免调用方省略
        # template_id 时因默认 editorial-modern 被拒绝。PDF/PPTX 承载完整主题差异。
        _template_id = options.template_id or TemplateId.editorial_modern

        locale = options.locale or "zh-CN"
        word_document = WordDocument()
        _configure_document(word_document, locale)
        word_document.core_properties.title = document.title
        word_document.core_properties.language = locale
        word_document.add_heading(document.title, level=0)
        if document.subtitle:
            subtitle = word_document.add_paragraph(style="Subtitle")
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            subtitle.add_run(document.subtitle)
        if options.include_toc:
            _add_table_of_contents(word_document, locale)

        for section in document.sections:
            word_document.add_heading(section.heading, level=section.level)
            for block in section.blocks:
                if isinstance(block, ParagraphBlock):
                    word_document.add_paragraph(block.text)
                elif isinstance(block, BulletListBlock):
                    for item in block.items:
                        word_document.add_paragraph(item, style="List Bullet")
                elif isinstance(block, NumberedListBlock):
                    for item in block.items:
                        word_document.add_paragraph(item, style="List Number")
                elif isinstance(block, TableBlock):
                    _add_table(word_document, block)
                elif isinstance(block, QuoteBlock):
                    quote_paragraph = word_document.add_paragraph(block.text, style="Quote")
                    if block.attribution:
                        attribution = word_document.add_paragraph()
                        attribution.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        attribution.add_run(block.attribution).italic = True
                    quote_paragraph.paragraph_format.keep_together = True
                elif isinstance(block, PageBreakBlock):
                    word_document.add_page_break()
                elif isinstance(block, ImageBlock):
                    _add_image(word_document, block)

        _add_page_numbers(word_document, locale)
        output = BytesIO()
        word_document.save(output)
        return RenderedDocx(
            content=output.getvalue(),
            filename=_safe_filename(document.title),
        )


def content_disposition(filename: str) -> str:
    encoded_filename = quote(filename, safe="")
    return f"attachment; filename=\"document.docx\"; filename*=UTF-8''{encoded_filename}"


def _configure_document(document: DocxDocument, locale: str) -> None:
    is_chinese = locale.lower().startswith("zh")
    body_font = "SimSun" if is_chinese else "Aptos"
    heading_font = "Microsoft YaHei" if is_chinese else "Aptos Display"

    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.18)
        section.right_margin = Cm(3.18)

    normal = document.styles["Normal"]
    _set_style_font(normal, body_font, 10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    heading_styles = (
        ("Title", 22),
        ("Subtitle", 12),
        ("Heading 1", 16),
        ("Heading 2", 14),
        ("Heading 3", 12),
    )
    for style_name, size in heading_styles:
        _set_style_font(document.styles[style_name], heading_font, size)


def _set_style_font(style: Any, font_name: str, size: float) -> None:
    style.font.name = font_name
    style.font.size = Pt(size)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)


def _add_image(document: DocxDocument, block: ImageBlock) -> None:
    """居中插入图片与图注；图片取不到时降级为一行说明，不中断整份文档。"""
    image = load_image(block.url)
    if image is None:
        label = block.alt or block.caption or "图片"
        placeholder = document.add_paragraph()
        placeholder.alignment = WD_ALIGN_PARAGRAPH.CENTER
        placeholder.add_run(f"[图片未能加载：{label}]").italic = True
        return

    width_cm = _CONTENT_WIDTH_CM * float(block.width_ratio or 1.0)
    height_cm = width_cm / image.aspect_ratio
    if height_cm > _MAX_IMAGE_HEIGHT_CM:
        height_cm = _MAX_IMAGE_HEIGHT_CM
        width_cm = height_cm * image.aspect_ratio

    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(
        BytesIO(image.content), width=Cm(width_cm), height=Cm(height_cm)
    )
    if block.caption:
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run(block.caption).italic = True


def _add_table(document: DocxDocument, block: TableBlock) -> None:
    table = document.add_table(rows=1, cols=len(block.columns))
    table.style = "Table Grid"
    for index, column in enumerate(block.columns):
        run = table.rows[0].cells[index].paragraphs[0].add_run(column)
        run.bold = True
    for values in block.rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value


def _add_table_of_contents(document: DocxDocument, locale: str) -> None:
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.add_run("目录" if locale.lower().startswith("zh") else "Table of Contents").bold = True
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, end))
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    document.settings.element.append(update_fields)
    document.add_page_break()


def _add_page_numbers(document: DocxDocument, locale: str) -> None:
    for section in document.sections:
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if locale.lower().startswith("zh"):
            paragraph.add_run("第 ")
        run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = " PAGE "
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.extend((begin, instruction, end))
        if locale.lower().startswith("zh"):
            paragraph.add_run(" 页")


def _safe_filename(title: str) -> str:
    filename = _INVALID_FILENAME.sub("_", title).strip(" .")
    if not filename:
        filename = "document"
    return f"{filename[:100].rstrip(' .') or 'document'}.docx"

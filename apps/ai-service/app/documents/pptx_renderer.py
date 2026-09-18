from __future__ import annotations

import math
import re
from dataclasses import dataclass
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from app.api.generated.models import (
    DocumentOptions,
    PptxBulletListBlock,
    PptxImageBlock,
    PptxNumberedListBlock,
    PptxParagraphBlock,
    PptxQuoteBlock,
    PptxSlide,
    PptxSpec,
    PptxTableBlock,
    Theme,
)
from app.documents.images import load_image
from app.documents.normalize import normalize_pptx_spec
from app.documents.text_metrics import looks_like_subheading
from app.documents.text_metrics import visual_units as _visual_units
from app.documents.validation import validate_pptx_spec

PPTX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
_INVALID_FILENAME = re.compile(r'[\x00-\x1f<>:"/\\|?*]+')

# PPTX 是 OOXML，字体由打开它的客户端渲染，无需内嵌；显式声明东亚字体保证中文跨平台。
_CJK_FONT = "Microsoft YaHei"

_BRAND = RGBColor(0x56, 0x5C, 0xF6)
_BRAND_SOFT = RGBColor(0xEE, 0xF0, 0xFF)
_TEXT = RGBColor(0x20, 0x24, 0x32)
_MUTED = RGBColor(0x85, 0x8C, 0x9B)
_LINE = RGBColor(0xE8, 0xEB, 0xF1)
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
_ROW_ALT = RGBColor(0xF7, 0xF8, 0xFB)

_NEUTRAL = RGBColor(0x5F, 0x67, 0x77)
_NEUTRAL_SOFT = RGBColor(0xF1, 0xF3, 0xF8)

# 16:9 版心几何（英寸）。所有版式坐标都由这些常量推导，避免散落的魔法数字。
_SLIDE_W = 13.333
_SLIDE_H = 7.5
_MARGIN_X = 0.72
_TITLE_LEFT = _MARGIN_X
_TITLE_TOP = 0.46
_TITLE_HEIGHT = 0.84
_RULE_TOP = 1.36
_RULE_WIDTH = 1.7
_RULE_HEIGHT = 0.055
_BODY_LEFT = _MARGIN_X
_BODY_TOP = 1.68
_BODY_WIDTH = _SLIDE_W - 2 * _MARGIN_X
_BODY_BOTTOM = 6.86
_FOOTER_RULE_TOP = 7.0
_FOOTER_TEXT_TOP = 7.06

_COVER_TITLE_SIZE = 40
_COVER_SUBTITLE_SIZE = 20
_SLIDE_TITLE_SIZE = 28
# 正文候选字号：幻灯片正文不应小于 14pt，否则投影/大屏上无法阅读。
# 内容放不下时优先「拆成更多页」，而不是把字号缩到看不清（旧实现会一路缩到 11pt，
# 结果是一页塞满小字的「文字墙」，是版式乱、样式丑的主因）。
_BODY_FONT_CANDIDATES = (20, 19, 18, 17, 16, 15, 14)
_MIN_BODY_FONT = _BODY_FONT_CANDIDATES[-1]

# 段落型小标题：短句、单行、不以句末标点结尾时按小标题排版，建立「标题 > 小标题 > 正文」层级。
# 小标题字号恒比正文大 2pt 且不低于 16pt，避免随正文字号缩小而一起塌掉层级。
_SUBHEADING_EXTRA_PT = 2
_MIN_SUBHEADING_SIZE = 16
# 小标题上方留白（按正文字号的换算比例），让它在视觉上归属下一段正文。
_SUBHEADING_GAP_RATIO = 0.45
# 小标题左侧品牌色竖条：给「标题 > 小标题 > 正文」再加一层可视线索，
# 否则单靠字号 +2pt 在投影上很难一眼分辨层级。
_SUBHEADING_BAR_WIDTH = 0.055
_SUBHEADING_INDENT = 0.2

# 行高系数：与渲染时写入 OOXML 的 spcPct 保持一致，让「自研度量」和「客户端真实
# 排版」同源。大字号（封面标题）在真实字体下的行盒明显更高，实测需要 1.42 才能
# 容下，否则文字会溢出框体压到下一块内容上。
_LINE_SPACING_BODY = 1.32
_LINE_SPACING_LARGE = 1.42
_LARGE_TEXT_PT = 24
_CELL_LINE_SPACING = 1.25

_BLOCK_GAP = 0.18
_LIST_INDENT = 0.34
_CELL_PAD_X = 0.09
_CELL_PAD_Y = 0.07
_IMAGE_BLOCK_HEIGHT = 3.2
_COLUMN_GUTTER = 0.5

# 续页标题后缀：一节内容超出单页时拆成多页，续页保留同标题并加此后缀。
_CONTINUATION_SUFFIX = "（续）"

# python-pptx 内置模板的「空白」版式索引。我们自绘标题/版心，不用模板占位符，
# 避免默认模板的空占位符与弱排版干扰版式。
_BLANK_LAYOUT_INDEX = 6

# 「无样式，表格网格」表样式：给出中性网格线，其余由我们自己着色。
_TABLE_GRID_STYLE_ID = "{5940675A-B579-460E-94D1-54222C63F5DA}"

# a:pPr 子元素的规范顺序，用于安全插入项目符号节点。
_PPR_ORDER = (
    "a:lnSpc",
    "a:spcBef",
    "a:spcAft",
    "a:buClrTx",
    "a:buClr",
    "a:buSzTx",
    "a:buSzPct",
    "a:buSzPts",
    "a:buFontTx",
    "a:buFont",
    "a:buNone",
    "a:buAutoNum",
    "a:buChar",
    "a:tabLst",
    "a:defRPr",
    "a:extLst",
)


@dataclass(frozen=True)
class RenderedPptx:
    content: bytes
    filename: str


class PptxRenderer:
    def render(
        self,
        spec: PptxSpec,
        options: DocumentOptions,
        *,
        request_id: str,
    ) -> RenderedPptx:
        # 渲染前统一清洗：合并 LLM 生成的「封面」小节、剥离「标题：」等标签前缀、
        # 剔除「讲师：____（占位）」占位行。否则会额外渲染一张重复且带占位符的封面页，
        # 是「版式乱、样式丑」的主要来源。
        spec = normalize_pptx_spec(spec)
        validate_pptx_spec(spec, request_id=request_id, status_code=422)
        theme = spec.theme or Theme.brand
        palette = _palette(theme)

        presentation = Presentation()
        presentation.slide_width = Inches(_SLIDE_W)
        presentation.slide_height = Inches(_SLIDE_H)

        # 先把结构化幻灯片按可用高度拆成实际页，得到准确总页数，避免页脚编号错位。
        pages = _plan_pages(spec.slides)

        _render_cover(presentation, spec, palette)
        total = len(pages)
        for index, page in enumerate(pages, start=1):
            _render_page(presentation, page, palette, spec.title, index, total)

        output = BytesIO()
        presentation.save(output)
        return RenderedPptx(
            content=output.getvalue(),
            filename=_safe_filename(spec.title),
        )


def _palette(theme: Theme) -> dict[str, RGBColor]:
    if theme == Theme.neutral:
        return {
            "primary": _NEUTRAL,
            "primary_soft": _NEUTRAL_SOFT,
            "text": _TEXT,
            "muted": _MUTED,
            "line": _LINE,
        }
    return {
        "primary": _BRAND,
        "primary_soft": _BRAND_SOFT,
        "text": _TEXT,
        "muted": _MUTED,
        "line": _LINE,
    }


def _render_cover(presentation: Presentation, spec: PptxSpec, palette: dict[str, RGBColor]) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[_BLANK_LAYOUT_INDEX])
    _set_background(slide, palette["primary"])

    # 顶部细白线作为品牌装饰，避免纯色底过于单调。
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(_MARGIN_X), Inches(2.5), Inches(2.2), Inches(_RULE_HEIGHT)
    )
    _style_shape(accent, fill=_WHITE)

    title_height = _text_box_height(
        _wrap_lines(spec.title, _COVER_TITLE_SIZE, _BODY_WIDTH), _COVER_TITLE_SIZE
    )
    title = slide.shapes.add_textbox(
        Inches(_MARGIN_X), Inches(2.72), Inches(_BODY_WIDTH), Inches(title_height)
    )
    _write_text(title, spec.title, _COVER_TITLE_SIZE, _WHITE, bold=True)

    if spec.subtitle:
        subtitle = slide.shapes.add_textbox(
            Inches(_MARGIN_X), Inches(2.72 + title_height + 0.16), Inches(_BODY_WIDTH), Inches(0.7)
        )
        _write_text(subtitle, spec.subtitle, _COVER_SUBTITLE_SIZE, _BRAND_SOFT)


@dataclass(frozen=True)
class _Page:
    """一页实际幻灯片：一个 section 可能拆成多页。"""

    title: str
    layout: str
    blocks: list
    font_pt: int
    continuation: bool


@dataclass(frozen=True)
class _ListChunk:
    """跨页拆开后的列表片段。

    列表是唯一允许跨页拆分的块：一个 11 条的目录清单若整块不可拆，要么把这一页
    撑到溢出，要么整块挪到下一页，在上一页留下大片空白（实测续页占用率不到 20%）。
    拆开时还要携带续页的起始编号，否则编号列表会在续页从 1 重新计数。
    """

    block: object
    start_at: int = 1


# 列表是唯一允许跨页拆分的块；拆分后每段不少于 2 条，避免出现只剩一条的碎片页。
_MIN_CHUNK_ITEMS = 2
_SPLITTABLE_BLOCKS = (PptxBulletListBlock, PptxNumberedListBlock)


def _plan_pages(slides: list[PptxSlide]) -> list[_Page]:
    """按版心可用高度把每个 section 拆成一页或多页。

    python-pptx 不会自动分页：内容超出版心时会直接溢出到幻灯片外并压到页脚上，
    这正是版面「乱」的根因。这里先测量每个块的高度，依次装箱，超出可用高度就
    另起一页，保证每页内容都落在版心内。
    """
    pages: list[_Page] = []
    for slide in slides:
        layout = _layout_name(slide.layout)
        if not slide.blocks or layout in ("title", "section_header"):
            pages.append(_Page(slide.title, layout, list(slide.blocks), _SLIDE_TITLE_SIZE, False))
            continue

        measure_width = _measure_width(layout)
        available = _BODY_BOTTOM - _BODY_TOP
        # 先尽量让整节落在一页；实在放不下时按最小字号分页，保证可读性下限。
        font_pt = _fit_font_size(slide.blocks, measure_width, available)
        for index, group in enumerate(_paginate(slide.blocks, measure_width, available, font_pt)):
            pages.append(_Page(slide.title, layout, group, font_pt, index > 0))
    return pages


def _measure_width(layout: str) -> float:
    if layout == "two_column":
        return (_BODY_WIDTH - _COLUMN_GUTTER) / 2
    return _BODY_WIDTH


def _is_splittable(block: object) -> bool:
    return isinstance(block, _SPLITTABLE_BLOCKS) and len(block.items) >= _MIN_CHUNK_ITEMS * 2


def _as_chunk(block: object, items: list[str], start_at: int) -> _ListChunk:
    """把列表的一段包装成续页片段，并把续页的起始编号带上。"""
    return _ListChunk(block=block.model_copy(update={"items": list(items)}), start_at=start_at)


def _split_list(
    block: object, font_pt: int, width: float, room: float
) -> tuple[_ListChunk, _ListChunk] | None:
    """把列表按剩余高度切成「当前页放得下的一段 + 续页的一段」。

    两段都不少于 2 条时才认为拆分有意义；否则返回 None，由调用方整块另起一页。
    """
    items = list(block.items)  # type: ignore[attr-defined]
    head_count = 0
    for count in range(1, len(items) + 1):
        if _measure_list(items[:count], font_pt, width) > room:
            break
        head_count = count
    if head_count < _MIN_CHUNK_ITEMS or len(items) - head_count < _MIN_CHUNK_ITEMS:
        return None
    return (
        _as_chunk(block, items[:head_count], 1),
        _as_chunk(block, items[head_count:], 1 + head_count),
    )


def _paginate(blocks: list, width: float, height: float, font_pt: int) -> list[list]:
    """按块高度把内容装箱到多页。

    列表按条目跨页拆分（编号连续），其余块整块移到下一页，避免把段落/表格切碎。
    """
    pages: list[list] = []
    current: list = []
    used = 0.0

    def flush() -> None:
        nonlocal current, used
        pages.append(current)
        current = []
        used = 0.0

    pending = list(blocks)
    while pending:
        block = pending.pop(0)
        block_height = _measure_block(block, font_pt, width)
        base = 0.0 if not current else used + _BLOCK_GAP
        if current and base + block_height > height:
            flush()
            base = 0.0
        # 放不下时尝试按条目拆分列表（含「单块本身就高于整页」的情况）。
        if _is_splittable(block) and block_height > height - base:
            split = _split_list(block, font_pt, width, height - base)
            if split is not None:
                head, tail = split
                current.append(head)
                used = base + _measure_block(head, font_pt, width)
                flush()
                pending.insert(0, tail)
                continue
            if current:
                flush()
                base = 0.0
        current.append(block)
        used = base + block_height
    if current:
        flush()
    return _avoid_orphan_subheadings(pages or [[]], width, height, font_pt)


def _is_subheading_block(block: object) -> bool:
    return isinstance(block, PptxParagraphBlock) and looks_like_subheading(block.text)


def _avoid_orphan_subheadings(
    pages: list[list], width: float, height: float, font_pt: int
) -> list[list]:
    """把落在页尾的小标题移到下一页，避免「小标题在页底、正文翻页」的孤立标题。

    只有当下一页仍容得下时才移动，否则宁可保留原地，也不制造新的溢出。
    """
    for index in range(len(pages) - 1):
        current, following = pages[index], pages[index + 1]
        moved: list = []
        while len(current) > 1 and _is_subheading_block(current[-1]):
            candidate = current[-1]
            if _measure_blocks(moved + [candidate] + following, font_pt, width) > height:
                break
            moved.insert(0, current.pop())
        following[:0] = moved
    return [page for page in pages if page] or [[]]


def _render_page(
    presentation: Presentation,
    page: _Page,
    palette: dict[str, RGBColor],
    doc_title: str,
    index: int,
    total: int,
) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[_BLANK_LAYOUT_INDEX])

    if page.layout == "section_header":
        _set_background(slide, palette["primary"])
        _render_centered_title(slide, page.title, _SLIDE_TITLE_SIZE + 6, _WHITE)
        _add_footer(slide, doc_title, index, total, palette, color=_BRAND_SOFT, rule=_WHITE)
        return

    if page.layout == "title":
        _render_centered_title(slide, page.title, _SLIDE_TITLE_SIZE, palette["text"])
        paragraphs = [block for block in page.blocks if isinstance(block, PptxParagraphBlock)]
        if paragraphs:
            subtitle = slide.shapes.add_textbox(
                Inches(_BODY_LEFT + 1.5), Inches(4.05), Inches(_BODY_WIDTH - 3.0), Inches(0.8)
            )
            _write_text(subtitle, paragraphs[0].text, 16, palette["muted"], align=PP_ALIGN.CENTER)
        _add_footer(slide, doc_title, index, total, palette)
        return

    title = f"{page.title}{_CONTINUATION_SUFFIX}" if page.continuation else page.title
    _render_slide_title(slide, title, palette)
    _add_footer(slide, doc_title, index, total, palette)

    available = _BODY_BOTTOM - _BODY_TOP
    if page.layout == "two_column":
        _render_two_column(slide, page.blocks, palette, available, page.font_pt)
    else:
        _render_flow(slide, page.blocks, palette, _BODY_TOP, _BODY_WIDTH, available, page.font_pt)


def _render_slide_title(slide, title: str, palette: dict[str, RGBColor]) -> None:
    box = slide.shapes.add_textbox(
        Inches(_TITLE_LEFT), Inches(_TITLE_TOP), Inches(_BODY_WIDTH), Inches(_TITLE_HEIGHT)
    )
    _write_text(box, title, _SLIDE_TITLE_SIZE, palette["primary"], bold=True)

    rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(_TITLE_LEFT),
        Inches(_RULE_TOP),
        Inches(_RULE_WIDTH),
        Inches(_RULE_HEIGHT),
    )
    _style_shape(rule, fill=palette["primary"])


def _render_centered_title(slide, title: str, size: int, color: RGBColor) -> None:
    height = _text_box_height(_wrap_lines(title, size, _BODY_WIDTH), size)
    box = slide.shapes.add_textbox(
        Inches(_BODY_LEFT), Inches(3.3), Inches(_BODY_WIDTH), Inches(height)
    )
    _write_text(box, title, size, color, bold=True, align=PP_ALIGN.CENTER)


def _add_footer(
    slide,
    doc_title: str,
    index: int,
    total: int,
    palette: dict[str, RGBColor],
    *,
    color: RGBColor | None = None,
    rule: RGBColor | None = None,
) -> None:
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(_MARGIN_X),
        Inches(_FOOTER_RULE_TOP),
        Inches(_BODY_WIDTH),
        Emu(9525),
    )
    _style_shape(line, fill=rule or palette["line"])

    label = slide.shapes.add_textbox(
        Inches(_MARGIN_X), Inches(_FOOTER_TEXT_TOP), Inches(_BODY_WIDTH - 1.4), Inches(0.3)
    )
    _write_text(label, doc_title, 9, color or palette["muted"])

    page = slide.shapes.add_textbox(
        Inches(_SLIDE_W - _MARGIN_X - 1.4), Inches(_FOOTER_TEXT_TOP), Inches(1.4), Inches(0.3)
    )
    _write_text(page, f"{index} / {total}", 9, color or palette["muted"], align=PP_ALIGN.RIGHT)


# --------------------------------------------------------------------------- #
# 内容流排版：以真实游标逐块下移，块之间不重叠。内容超出可用高度时按字号阶梯缩小。
# --------------------------------------------------------------------------- #


def _render_flow(
    slide,
    blocks: list,
    palette: dict[str, RGBColor],
    top: float,
    width: float,
    height: float,
    font_pt: int,
) -> None:
    if not blocks:
        return
    cursor = top
    for block in blocks:
        remaining = max(0.5, height - (cursor - top))
        cursor = _render_block(slide, block, palette, _BODY_LEFT, cursor, width, remaining, font_pt)
        cursor += _BLOCK_GAP


def _render_two_column(
    slide, blocks: list, palette: dict[str, RGBColor], height: float, font_pt: int
) -> None:
    if not blocks:
        return
    column_width = (_BODY_WIDTH - _COLUMN_GUTTER) / 2
    left_blocks, right_blocks = _split_columns(blocks, column_width)

    for column_index, column_blocks in enumerate((left_blocks, right_blocks)):
        if not column_blocks:
            continue
        left = _BODY_LEFT + column_index * (column_width + _COLUMN_GUTTER)
        cursor = _BODY_TOP
        for block in column_blocks:
            remaining = max(0.5, height - (cursor - _BODY_TOP))
            cursor = _render_block(
                slide, block, palette, left, cursor, column_width, remaining, font_pt
            )
            cursor += _BLOCK_GAP


def _split_columns(blocks: list, width: float) -> tuple[list, list]:
    """按累计高度把块切成左右两列，尽量让两列等高。"""
    total = sum(_measure_block(block, _MIN_BODY_FONT, width) for block in blocks)
    target = total / 2
    left: list = []
    running = 0.0
    index = 0
    while index < len(blocks) - 1:
        cost = _measure_block(blocks[index], _MIN_BODY_FONT, width)
        if left and running + cost > target:
            break
        left.append(blocks[index])
        running += cost
        index += 1
    if not left:
        left = [blocks[0]]
        index = 1
    return left, blocks[index:]


def _render_block(
    slide,
    block: object,
    palette: dict[str, RGBColor],
    left: float,
    top: float,
    width: float,
    height: float,
    font_pt: int,
) -> float:
    """渲染单个块并返回其底部坐标（下一次排版的游标）。"""
    if isinstance(block, PptxParagraphBlock):
        if looks_like_subheading(block.text):
            return _render_subheading(slide, block.text, palette, left, top, width, font_pt)
        box_height = _text_box_height(_wrap_lines(block.text, font_pt, width), font_pt)
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(box_height))
        _write_text(box, block.text, font_pt, palette["text"])
        return top + box_height

    if isinstance(block, _ListChunk):
        return _render_list(
            slide,
            list(block.block.items),  # type: ignore[attr-defined]
            palette,
            left,
            top,
            width,
            font_pt,
            bullet=isinstance(block.block, PptxBulletListBlock),
            start_at=block.start_at,
        )

    if isinstance(block, PptxBulletListBlock):
        return _render_list(slide, block.items, palette, left, top, width, font_pt, bullet=True)

    if isinstance(block, PptxNumberedListBlock):
        return _render_list(slide, block.items, palette, left, top, width, font_pt, bullet=False)

    if isinstance(block, PptxQuoteBlock):
        return _render_quote(slide, block, palette, left, top, width, font_pt)

    if isinstance(block, PptxTableBlock):
        return top + _render_table(slide, block, palette, left, top, width, font_pt)

    if isinstance(block, PptxImageBlock):
        return _render_image(slide, block, palette, left, top, width, height)

    return top


def _subheading_size(font_pt: int) -> int:
    """小标题字号：比正文大 2pt，且不低于 16pt、不超过幻灯片标题。"""
    return max(_MIN_SUBHEADING_SIZE, min(_SLIDE_TITLE_SIZE - 4, font_pt + _SUBHEADING_EXTRA_PT))


def _render_subheading(
    slide,
    text: str,
    palette: dict[str, RGBColor],
    left: float,
    top: float,
    width: float,
    font_pt: int,
) -> float:
    """按小标题排版一段短句：品牌色加粗、左侧色条、字号更大、上方留白。

    LLM 生成的讲义常把节内小标题写成普通段落块，若与正文同字号同颜色排下来，
    整页就是一堵没有层级的文字墙。这里把「标题 > 小标题 > 正文」的层级显式画出来。
    """
    size = _subheading_size(font_pt)
    gap = _SUBHEADING_GAP_RATIO * font_pt / 72.0
    top += gap
    text_left = left + _SUBHEADING_INDENT
    text_width = width - _SUBHEADING_INDENT
    box_height = _text_box_height(_wrap_lines(text, size, text_width), size)

    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left),
        Inches(top + 0.01),
        Inches(_SUBHEADING_BAR_WIDTH),
        Inches(max(0.14, box_height - 0.02)),
    )
    _style_shape(bar, fill=palette["primary"])

    box = slide.shapes.add_textbox(
        Inches(text_left), Inches(top), Inches(text_width), Inches(box_height)
    )
    _write_text(box, text, size, palette["primary"], bold=True)
    return top + box_height


def _render_list(
    slide,
    items: list[str],
    palette: dict[str, RGBColor],
    left: float,
    top: float,
    width: float,
    font_pt: int,
    *,
    bullet: bool,
    start_at: int = 1,
) -> float:
    total = _measure_list(items, font_pt, width)
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(total))
    frame = _prepare_frame(box)
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = item
        _style_run(run, font_pt, palette["text"], False)
        if bullet:
            _set_bullet(paragraph, char="•", marL=_LIST_INDENT)
        else:
            _set_bullet(
                paragraph,
                auto_num="arabicPeriod",
                marL=_LIST_INDENT,
                auto_num_start=start_at,
            )
        # 行距必须写进 XML，否则客户端默认行距高于我们测量的行高，文字会溢出框体。
        _set_line_spacing(paragraph, _LINE_SPACING_BODY)
        # 段后留白：让多行条目有可读的间隙。
        _set_paragraph_spacing(paragraph, space_after_pt=font_pt * 0.4)
    return top + total


def _render_quote(
    slide,
    block: PptxQuoteBlock,
    palette: dict[str, RGBColor],
    left: float,
    top: float,
    width: float,
    font_pt: int,
) -> float:
    size = max(11, font_pt - 1)
    inner_width = width - 0.7
    lines = _wrap_lines(block.text, size, inner_width) + (1 if block.attribution else 0)
    pad = 0.2
    height = _text_box_height(lines, size) + pad * 2

    background = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    _style_shape(background, fill=palette["primary_soft"])
    background.adjustments[0] = 0.06

    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(0.06), Inches(height)
    )
    _style_shape(bar, fill=palette["primary"])

    box = slide.shapes.add_textbox(
        Inches(left + 0.34), Inches(top + pad), Inches(inner_width), Inches(height - pad * 2)
    )
    frame = _prepare_frame(box)
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = block.text
    _style_run(run, size, palette["text"], False)
    _set_line_spacing(paragraph, _LINE_SPACING_BODY)
    if block.attribution:
        attribution = frame.add_paragraph()
        attribution.alignment = PP_ALIGN.RIGHT
        run = attribution.add_run()
        run.text = f"— {block.attribution}"
        _style_run(run, max(10, size - 1), palette["muted"], False)
        _set_line_spacing(attribution, _LINE_SPACING_BODY)
    return top + height


def _render_table(
    slide,
    block: PptxTableBlock,
    palette: dict[str, RGBColor],
    left: float,
    top: float,
    width: float,
    font_pt: int,
) -> float:
    cols = len(block.columns)
    rows = len(block.rows) + 1
    size = max(11, font_pt - 2)
    column_widths = _column_widths(block, width)
    row_heights = _row_heights(block, column_widths, size)
    total_height = sum(row_heights)

    graphic_frame = slide.shapes.add_table(
        rows, cols, Inches(left), Inches(top), Inches(width), Inches(total_height)
    )
    table = graphic_frame.table
    table.first_row = False
    table.horz_banding = False
    _set_table_style(graphic_frame, _TABLE_GRID_STYLE_ID)

    for index, column_width in enumerate(column_widths):
        table.columns[index].width = Inches(column_width)
    for index, row_height in enumerate(row_heights):
        table.rows[index].height = Inches(row_height)

    for col_index, label in enumerate(block.columns):
        _fill_cell(
            table.cell(0, col_index), label, size, _WHITE, bold=True, fill=palette["primary"]
        )
    for row_index, values in enumerate(block.rows, start=1):
        fill = _neutral_row_fill(palette, _ROW_ALT if row_index % 2 == 0 else _WHITE)
        for col_index in range(cols):
            value = values[col_index] if col_index < len(values) else ""
            _fill_cell(
                table.cell(row_index, col_index),
                value,
                size,
                palette["text"],
                bold=False,
                fill=fill,
            )
    return total_height


def _render_image(
    slide,
    block: PptxImageBlock,
    palette: dict[str, RGBColor],
    left: float,
    top: float,
    width: float,
    height: float,
) -> float:
    max_height = min(height, _IMAGE_BLOCK_HEIGHT + 0.8)
    # 与 DOCX/PDF 共用同一取图逻辑：只接受 http(s) 与 data URL、限制体积，
    # 失败时返回 None 并降级为占位文本，绝不让一张失效的图打断整页渲染。
    image = load_image(block.url)
    if image is None:
        label = block.alt or block.caption or block.url
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.6))
        _write_text(box, f"[图片未能加载：{label}]", 14, palette["muted"])
        return top + 0.6

    shown_width = min(width, 8.5)
    shown_height = shown_width / image.aspect_ratio
    if shown_height > max_height:
        shown_height = max_height
        shown_width = shown_height * image.aspect_ratio
    offset_x = left + (width - shown_width) / 2
    slide.shapes.add_picture(
        BytesIO(image.content),
        Inches(offset_x),
        Inches(top),
        width=Inches(shown_width),
        height=Inches(shown_height),
    )
    bottom = top + shown_height
    if block.caption:
        box = slide.shapes.add_textbox(
            Inches(left), Inches(bottom + 0.08), Inches(width), Inches(0.4)
        )
        _write_text(box, block.caption, 11, palette["muted"], align=PP_ALIGN.CENTER)
        bottom += 0.48
    return bottom


# --------------------------------------------------------------------------- #
# 度量：块高度估算。python-pptx 不提供文本测量，用「中日韩字符=1 个字宽、
# 其它字符≈0.55 个字宽」的近似换算成行数，再乘以行高，得到稳定的排版游标。
# --------------------------------------------------------------------------- #


def _line_spacing_for(font_pt: float) -> float:
    """按字号选择行距系数。

    大字号（封面标题、幻灯片标题）在真实字体下的行盒明显高于小字号，用同一个
    系数会让大字号文本框被算矮、文字溢出压到下一块内容上。这里以 _LARGE_TEXT_PT
    为界切换系数，并保证「度量」与「写入 OOXML 的 spcPct」同源。
    """
    return _LINE_SPACING_LARGE if font_pt >= _LARGE_TEXT_PT else _LINE_SPACING_BODY


def _wrap_lines(text: str, font_pt: float, width_in: float) -> int:
    per_line = max(1.0, max(1.0, width_in * 72.0) / float(font_pt))
    lines = 0
    for part in str(text).split("\n"):
        lines += max(1, math.ceil(_visual_units(part) / per_line))
    return max(1, lines)


def _text_box_height(
    lines: int, font_pt: float, *, line_spacing: float | None = None, pad: float = 0.0
) -> float:
    """按行数估算文本框高度。

    上下留白随字号增大：真实字体（微软雅黑等）在大字号下的行高比线性估算更高，
    固定 0.06in 的留白在 40pt 封面标题上会不够（实测需要多约 12% 高度）。
    """
    spacing = _line_spacing_for(font_pt) if line_spacing is None else line_spacing
    effective_pad = pad or (0.06 + font_pt / 72.0 * 0.10)
    return lines * font_pt * spacing / 72.0 + effective_pad


def _measure_list(items: list[str], font_pt: int, width: float) -> float:
    text_width = width - _LIST_INDENT
    total = 0.0
    for item in items:
        total += _text_box_height(_wrap_lines(item, font_pt, text_width), font_pt)
        total += font_pt * 0.4 / 72.0
    return max(total, _text_box_height(1, font_pt))


def _measure_block(block: object, font_pt: int, width: float) -> float:
    if isinstance(block, PptxParagraphBlock):
        if looks_like_subheading(block.text):
            size = _subheading_size(font_pt)
            gap = _SUBHEADING_GAP_RATIO * font_pt / 72.0
            text_width = width - _SUBHEADING_INDENT
            return gap + _text_box_height(_wrap_lines(block.text, size, text_width), size)
        return _text_box_height(_wrap_lines(block.text, font_pt, width), font_pt)
    if isinstance(block, _ListChunk):
        return _measure_list(list(block.block.items), font_pt, width)  # type: ignore[attr-defined]
    if isinstance(block, (PptxBulletListBlock, PptxNumberedListBlock)):
        return _measure_list(block.items, font_pt, width)
    if isinstance(block, PptxQuoteBlock):
        size = max(11, font_pt - 1)
        lines = _wrap_lines(block.text, size, width - 0.7) + (1 if block.attribution else 0)
        return _text_box_height(lines, size) + 0.4
    if isinstance(block, PptxTableBlock):
        return sum(_row_heights(block, _column_widths(block, width), max(11, font_pt - 2)))
    if isinstance(block, PptxImageBlock):
        return _IMAGE_BLOCK_HEIGHT
    return 0.3


def _measure_blocks(blocks: list, font_pt: int, width: float) -> float:
    if not blocks:
        return 0.0
    total = sum(_measure_block(block, font_pt, width) for block in blocks)
    return total + _BLOCK_GAP * (len(blocks) - 1)


def _fit_font_size(blocks: list, width: float, height: float) -> int:
    for size in _BODY_FONT_CANDIDATES:
        if _measure_blocks(blocks, size, width) <= height:
            return size
    return _MIN_BODY_FONT


def _column_widths(block: PptxTableBlock, total_width: float) -> list[float]:
    cols = len(block.columns)
    weights: list[float] = []
    for col_index in range(cols):
        weight = _visual_units(block.columns[col_index])
        for row in block.rows:
            if col_index < len(row):
                weight = max(weight, _visual_units(row[col_index]))
        weights.append(max(1.0, weight))

    minimum = min(1.1, total_width / cols)
    provisional = [max(minimum, total_width * weight / sum(weights)) for weight in weights]
    scale = total_width / sum(provisional)
    return [width * scale for width in provisional]


def _row_heights(block: PptxTableBlock, column_widths: list[float], font_pt: int) -> list[float]:
    line_height = font_pt * _CELL_LINE_SPACING / 72.0
    padding = _CELL_PAD_Y * 2

    def row_lines(values: list[str]) -> int:
        lines = 1
        for index, width in enumerate(column_widths):
            value = values[index] if index < len(values) else ""
            lines = max(lines, _wrap_lines(value, font_pt, max(0.4, width - _CELL_PAD_X * 2)))
        return lines

    heights = [row_lines(list(block.columns)) * line_height + padding]
    for row in block.rows:
        heights.append(row_lines(list(row)) * line_height + padding)
    return heights


def _prepare_frame(box) -> object:
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    _set_shrink_on_overflow(frame)
    return frame


def _set_shrink_on_overflow(frame) -> None:
    """为文本框声明「溢出时自动缩字」（a:normAutofit）。

    排版游标来自自研度量，与客户端真实字体度量存在小偏差；一旦偏差为「实际更高」，
    文字就会溢出框体并压到下一块内容上（最典型的「版式乱」）。声明 normAutofit 后，
    这种偏差由客户端缩字兜底，而不是靠估算做到零误差。
    """
    body_pr = frame._txBody.find(qn("a:bodyPr"))
    if body_pr is None:
        return
    for tag in ("a:noAutofit", "a:normAutofit", "a:spAutoFit"):
        for existing in body_pr.findall(qn(tag)):
            body_pr.remove(existing)
    element = body_pr.makeelement(qn("a:normAutofit"), {})
    # a:bodyPr 子元素有固定顺序：autofit 必须排在 extLst 之前。
    ext_lst = body_pr.find(qn("a:extLst"))
    if ext_lst is not None:
        ext_lst.addprevious(element)
    else:
        body_pr.append(element)


def _write_text(
    box,
    text: str,
    size: int,
    color: RGBColor,
    *,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    frame = _prepare_frame(box)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    _style_run(run, size, color, bold)
    # 行距与度量同源（见 _line_spacing_for），否则大字号标题会溢出框体。
    _set_line_spacing(paragraph, _line_spacing_for(size))


def _style_run(run, size: float, color: RGBColor, bold: bool) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = _CJK_FONT
    run.font.color.rgb = color
    # 显式设置 latin/eastAsia/cs 字体，保证中文跨平台正常渲染
    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        existing = r_pr.find(qn(tag))
        if existing is not None:
            r_pr.remove(existing)
        element = r_pr.makeelement(qn(tag), {"typeface": _CJK_FONT})
        r_pr.append(element)


def _set_line_spacing(paragraph, spacing: float) -> None:
    """写入行距百分比（a:lnSpc/a:spcPct）。

    与 _text_box_height 使用同一个系数，才能保证「自研度量」和「客户端真实排版」
    一致；否则默认行距会把多行文本撑出框体、压到下一块内容上。
    """
    p_pr = paragraph._p.get_or_add_pPr()
    element = p_pr.makeelement(qn("a:lnSpc"), {})
    element.append(element.makeelement(qn("a:spcPct"), {"val": str(int(spacing * 100000))}))
    _insert_ppr_child(p_pr, element, "a:lnSpc")


def _set_paragraph_spacing(paragraph, *, space_after_pt: float) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    element = p_pr.makeelement(qn("a:spcAft"), {})
    element.append(element.makeelement(qn("a:spcPts"), {"val": str(int(space_after_pt * 100))}))
    _insert_ppr_child(p_pr, element, "a:spcAft")


def _set_bullet(
    paragraph,
    *,
    char: str | None = None,
    auto_num: str | None = None,
    marL: float,
    auto_num_start: int = 1,
) -> None:
    emu = int(marL * 914400)
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("marL", str(emu))
    p_pr.set("indent", str(-emu))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        existing = p_pr.find(qn(tag))
        if existing is not None:
            p_pr.remove(existing)
    if auto_num is not None:
        element = p_pr.makeelement(qn("a:buAutoNum"), {"type": auto_num})
        # 列表跨页拆分后续页必须从原编号继续，否则会从 1 重新计数。
        if auto_num_start > 1:
            element.set("startAt", str(auto_num_start))
        _insert_ppr_child(p_pr, element, "a:buAutoNum")
    elif char is not None:
        element = p_pr.makeelement(qn("a:buChar"), {"char": char})
        _insert_ppr_child(p_pr, element, "a:buChar")


def _insert_ppr_child(p_pr, element, tag: str) -> None:
    """按 a:pPr 的规范子元素顺序插入节点，避免生成 PowerPoint 拒绝的非法顺序。"""
    position = _PPR_ORDER.index(tag)
    for child in p_pr:
        qualified = f"a:{child.tag.split('}')[-1]}"
        if qualified in _PPR_ORDER and _PPR_ORDER.index(qualified) > position:
            child.addprevious(element)
            return
    p_pr.append(element)


def _fill_cell(
    cell,
    text: str,
    size: int,
    color: RGBColor,
    *,
    bold: bool,
    fill: RGBColor,
) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = Inches(_CELL_PAD_X)
    cell.margin_right = Inches(_CELL_PAD_X)
    cell.margin_top = Inches(_CELL_PAD_Y / 2)
    cell.margin_bottom = Inches(_CELL_PAD_Y / 2)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    frame = cell.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = text
    _style_run(run, size, color, bold)
    # 单元格行距与 _row_heights 的估算同源，避免行高算矮、文字被裁切。
    _set_line_spacing(paragraph, _CELL_LINE_SPACING)


def _neutral_row_fill(palette: dict[str, RGBColor], fill: RGBColor) -> RGBColor:
    if palette["primary"] == _NEUTRAL:
        return _NEUTRAL_SOFT if fill == _ROW_ALT else _WHITE
    return fill


def _style_shape(
    shape,
    *,
    fill: RGBColor | None = None,
    line_color: RGBColor | None = None,
    line_width_pt: float = 1.0,
) -> None:
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_width_pt)
    try:
        shape.shadow.inherit = False
    except (AttributeError, NotImplementedError):  # pragma: no cover - 版本兼容
        pass


def _set_background(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _set_table_style(graphic_frame, style_id: str) -> None:
    tbl_pr = graphic_frame.table._tbl.tblPr
    for existing in tbl_pr.findall(qn("a:tableStyleId")):
        tbl_pr.remove(existing)
    element = tbl_pr.makeelement(qn("a:tableStyleId"), {})
    element.text = style_id
    tbl_pr.append(element)


def _layout_name(layout: object) -> str:
    return str(getattr(layout, "value", layout))


def _safe_filename(title: str) -> str:
    filename = _INVALID_FILENAME.sub("_", title).strip(" .")
    if not filename:
        filename = "presentation"
    return f"{filename[:100].rstrip(' .') or 'presentation'}.pptx"
